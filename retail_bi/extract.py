"""Read UCI XLSX without modifying it; retain source values in CSV and controls."""
from collections import Counter
from datetime import datetime
from decimal import Decimal
from pathlib import Path
import csv
import hashlib
import json
import zipfile
from openpyxl import load_workbook
from .rules import classify

ROOT = Path(__file__).resolve().parents[1]
HEADERS = ['InvoiceNo', 'StockCode', 'Description', 'Quantity', 'InvoiceDate', 'UnitPrice', 'CustomerID', 'Country']


def value_text(value):
    if value is None:
        return ''
    if isinstance(value, datetime):
        return value.isoformat(sep=' ')
    return str(value)


def main():
    archive = ROOT/'data/raw/online-retail.zip'
    with zipfile.ZipFile(archive) as z:
        names = [n for n in z.namelist() if n.lower().endswith('.xlsx')]
        if len(names) != 1:
            raise ValueError('Expected exactly one XLSX in the UCI archive')
        target = ROOT/'data/raw/Online Retail.xlsx'
        target.write_bytes(z.read(names[0]))
    digest = hashlib.sha256(target.read_bytes()).hexdigest()
    manifest=json.loads((ROOT/'source-manifest.json').read_text())
    if digest!=manifest['xlsx_sha256']:raise ValueError('Source XLSX checksum differs from the pinned UCI version')
    (ROOT/'data/interim').mkdir(parents=True,exist_ok=True)
    (ROOT/'artifacts').mkdir(parents=True,exist_ok=True)
    book = load_workbook(target, read_only=True, data_only=True)
    iterator = book.active.iter_rows(values_only=True)
    if list(next(iterator)) != HEADERS:
        raise ValueError('Source column contract differs from UCI Online Retail')
    counts=Counter(); seen=set(); invoices=set(); customers=set()
    gross=Decimal(0); adjustments=Decimal(0); min_date=None; max_date=None
    output=ROOT/'data/interim/online-retail.csv'
    with output.open('w',newline='',encoding='utf-8') as f:
        writer=csv.writer(f); writer.writerow(['source_row',*HEADERS])
        for rownum,raw in enumerate(iterator,start=2):
            row=tuple(value_text(v) for v in raw)
            writer.writerow([rownum,*row]); counts['source_rows']+=1
            if not isinstance(raw[4],datetime):
                raise ValueError(f'Invalid source date at row {rownum}')
            min_date=raw[4] if min_date is None else min(min_date,raw[4])
            max_date=raw[4] if max_date is None else max(max_date,raw[4])
            if raw[6] is None:counts['missing_customer_rows']+=1
            if raw[2] is None:counts['missing_description_rows']+=1
            if row in seen:counts['duplicate_rows']+=1;continue
            seen.add(row)
            kind=classify(row[0],row[1],int(raw[3]),Decimal(str(raw[5])))
            counts[kind]+=1
            amount=Decimal(str(raw[3]))*Decimal(str(raw[5]))
            if kind=='sale':
                gross+=amount;invoices.add(row[0])
                if row[6]:customers.add(row[6])
            elif kind=='negative_adjustment':adjustments+=amount
    book.close()
    if counts['source_rows'] != 541909:
        raise ValueError('UCI source row count changed; inspect before accepting')
    profile={
        'source':'https://archive.ics.uci.edu/dataset/352/online+retail',
        'download_url':'https://archive.ics.uci.edu/static/public/352/online%2Bretail.zip',
        'citation':'Chen, D. (2015). Online Retail. UCI. DOI: 10.24432/C5BW33',
        'license':'CC BY 4.0','xlsx_sha256':digest,
        'counts':dict(counts),'gross_sales_gbp':str(gross),'negative_adjustments_gbp':str(adjustments),
        'sales_invoices':len(invoices),'known_customers':len(customers),
        'min_timestamp':min_date.isoformat(),'max_timestamp':max_date.isoformat(),
        'headers':HEADERS,'csv_sha256':hashlib.sha256(output.read_bytes()).hexdigest(),
    }
    (ROOT/'artifacts/source-profile.json').write_text(json.dumps(profile,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({k:profile[k] for k in ['counts','gross_sales_gbp','negative_adjustments_gbp','sales_invoices','known_customers','min_timestamp','max_timestamp']},ensure_ascii=False))


if __name__=='__main__':main()
